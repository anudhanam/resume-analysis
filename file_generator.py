"""
file_generator.py
Generate PDF and DOCX files from optimized resume data.
"""

import logging
import os
import tempfile
from typing import Any
from datetime import datetime

logger = logging.getLogger(__name__)


def generate_docx(resume_data: dict[str, Any], output_path: str = None) -> str:
    """
    Generate a DOCX resume from structured data.

    Args:
        resume_data: Structured resume dictionary.
        output_path: Optional output file path. If None, creates temp file.

    Returns:
        Path to generated DOCX file.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip install python-docx")

    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Helper functions
    def add_heading(text: str, level: int = 1):
        para = doc.add_paragraph()
        run = para.add_run(text)
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
        else:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x16, 0x21, 0x3e)
            para.paragraph_format.space_before = Pt(8)
            # Add bottom border
            pPr = para._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '162139')
            pBdr.append(bottom)
            pPr.append(pBdr)
        return para

    def add_contact_line(text: str):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        return para

    def add_body_text(text: str, bold: bool = False):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.size = Pt(10)
        run.font.bold = bold
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(1)
        return para

    def add_bullet(text: str):
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(text)
        run.font.size = Pt(10)
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(1)
        return para

    # ── NAME & CONTACT ──────────────────────────────────────────
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(resume_data.get("name", "Your Name"))
    name_run.font.size = Pt(20)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
    name_para.paragraph_format.space_after = Pt(2)

    contact_parts = []
    if resume_data.get("email"):
        contact_parts.append(resume_data["email"])
    if resume_data.get("phone"):
        contact_parts.append(resume_data["phone"])
    if contact_parts:
        add_contact_line(" | ".join(contact_parts))

    doc.add_paragraph()

    # ── PROFESSIONAL SUMMARY ────────────────────────────────────
    summary = resume_data.get("summary", "")
    if summary:
        add_heading("PROFESSIONAL SUMMARY", level=2)
        add_body_text(summary)

    # ── SKILLS ──────────────────────────────────────────────────
    all_skills = (
        resume_data.get("skills", []) +
        resume_data.get("tools", []) +
        resume_data.get("frameworks", [])
    )
    all_skills = list(dict.fromkeys(s for s in all_skills if s))  # deduplicate

    if all_skills:
        add_heading("TECHNICAL SKILLS", level=2)
        skills_text = " • ".join(all_skills)
        add_body_text(skills_text)

    # ── EXPERIENCE ──────────────────────────────────────────────
    experience = resume_data.get("experience", [])
    if experience:
        add_heading("PROFESSIONAL EXPERIENCE", level=2)
        for job in experience:
            role = job.get("role", "")
            company = job.get("company", "")
            duration = job.get("duration", "")
            header = f"{role} — {company}"
            if duration:
                header += f"  |  {duration}"
            add_body_text(header, bold=True)
            for bullet in job.get("bullets", []):
                if bullet:
                    add_bullet(bullet)
            doc.add_paragraph()

    # ── PROJECTS ────────────────────────────────────────────────
    projects = resume_data.get("projects", [])
    if projects:
        add_heading("PROJECTS", level=2)
        for project in projects:
            proj_name = project.get("name", "")
            description = project.get("description", "")
            techs = project.get("technologies", [])

            proj_para = doc.add_paragraph()
            run = proj_para.add_run(proj_name)
            run.font.bold = True
            run.font.size = Pt(10)

            if description:
                add_body_text(description)
            if techs:
                add_body_text(f"Technologies: {', '.join(techs)}")
            doc.add_paragraph()

    # ── EDUCATION ───────────────────────────────────────────────
    education = resume_data.get("education", [])
    if education:
        add_heading("EDUCATION", level=2)
        for edu in education:
            degree = edu.get("degree", "")
            institution = edu.get("institution", "")
            year = edu.get("year", "")
            edu_text = f"{degree} — {institution}"
            if year:
                edu_text += f"  |  {year}"
            add_body_text(edu_text, bold=False)

    # ── CERTIFICATIONS ──────────────────────────────────────────
    certs = resume_data.get("certifications", [])
    if certs:
        add_heading("CERTIFICATIONS", level=2)
        for cert in certs:
            if cert:
                add_bullet(cert)

    # Save
    if not output_path:
        tmp = tempfile.NamedTemporaryFile(
            delete=False, suffix=".docx",
            prefix=f"optimized_resume_{datetime.now().strftime('%Y%m%d_%H%M%S')}_"
        )
        output_path = tmp.name
        tmp.close()

    doc.save(output_path)
    logger.info(f"DOCX saved to: {output_path}")
    return output_path


def generate_pdf(resume_data: dict[str, Any], output_path: str = None) -> str:
    """
    Generate a PDF resume from structured data using reportlab.

    Args:
        resume_data: Structured resume dictionary.
        output_path: Optional output file path. If None, creates temp file.

    Returns:
        Path to generated PDF file.
    """
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer,
            HRFlowable, ListFlowable, ListItem
        )
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
    except ImportError:
        raise ImportError("reportlab not installed. Run: pip install reportlab")

    if not output_path:
        tmp = tempfile.NamedTemporaryFile(
            delete=False, suffix=".pdf",
            prefix=f"optimized_resume_{datetime.now().strftime('%Y%m%d_%H%M%S')}_"
        )
        output_path = tmp.name
        tmp.close()

    # Document setup
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
    )

    # Color palette
    dark_navy = colors.HexColor('#1a1a2e')
    medium_navy = colors.HexColor('#16213e')
    accent_blue = colors.HexColor('#0f3460')
    gray = colors.HexColor('#555555')
    light_gray = colors.HexColor('#888888')

    # Styles
    styles = {
        "name": ParagraphStyle(
            "name", fontName="Helvetica-Bold",
            fontSize=20, textColor=dark_navy,
            alignment=TA_CENTER, spaceAfter=4,
        ),
        "contact": ParagraphStyle(
            "contact", fontName="Helvetica",
            fontSize=9.5, textColor=gray,
            alignment=TA_CENTER, spaceAfter=2,
        ),
        "section_heading": ParagraphStyle(
            "section_heading", fontName="Helvetica-Bold",
            fontSize=11, textColor=medium_navy,
            spaceBefore=10, spaceAfter=4,
            borderPadding=(0, 0, 2, 0),
        ),
        "body": ParagraphStyle(
            "body", fontName="Helvetica",
            fontSize=9.5, textColor=colors.black,
            leading=14, spaceAfter=3,
        ),
        "bold_body": ParagraphStyle(
            "bold_body", fontName="Helvetica-Bold",
            fontSize=9.5, textColor=colors.black,
            leading=14, spaceAfter=2,
        ),
        "bullet": ParagraphStyle(
            "bullet", fontName="Helvetica",
            fontSize=9.5, textColor=colors.black,
            leading=14, leftIndent=12, spaceAfter=2,
            bulletIndent=0,
        ),
    }

    story = []

    def add_section(title: str):
        story.append(Spacer(1, 6))
        story.append(Paragraph(title, styles["section_heading"]))
        story.append(HRFlowable(
            width="100%", thickness=1.2,
            color=accent_blue, spaceAfter=4,
        ))

    # ── NAME & CONTACT ──────────────────────────────────────────
    story.append(Paragraph(
        resume_data.get("name", "Your Name"),
        styles["name"]
    ))

    contact_parts = []
    if resume_data.get("email"):
        contact_parts.append(resume_data["email"])
    if resume_data.get("phone"):
        contact_parts.append(resume_data["phone"])
    if contact_parts:
        story.append(Paragraph(" | ".join(contact_parts), styles["contact"]))

    story.append(Spacer(1, 6))

    # ── SUMMARY ─────────────────────────────────────────────────
    if resume_data.get("summary"):
        add_section("PROFESSIONAL SUMMARY")
        story.append(Paragraph(resume_data["summary"], styles["body"]))

    # ── SKILLS ──────────────────────────────────────────────────
    all_skills = list(dict.fromkeys(
        s for s in (
            resume_data.get("skills", []) +
            resume_data.get("tools", []) +
            resume_data.get("frameworks", [])
        ) if s
    ))
    if all_skills:
        add_section("TECHNICAL SKILLS")
        story.append(Paragraph(" • ".join(all_skills), styles["body"]))

    # ── EXPERIENCE ──────────────────────────────────────────────
    experience = resume_data.get("experience", [])
    if experience:
        add_section("PROFESSIONAL EXPERIENCE")
        for job in experience:
            role = job.get("role", "")
            company = job.get("company", "")
            duration = job.get("duration", "")
            header = f"<b>{role}</b> — {company}"
            if duration:
                header += f"  <font color='#888888' size='9'>{duration}</font>"
            story.append(Paragraph(header, styles["bold_body"]))
            for bullet in job.get("bullets", []):
                if bullet:
                    story.append(Paragraph(f"• {bullet}", styles["bullet"]))
            story.append(Spacer(1, 4))

    # ── PROJECTS ────────────────────────────────────────────────
    projects = resume_data.get("projects", [])
    if projects:
        add_section("PROJECTS")
        for project in projects:
            name = project.get("name", "")
            desc = project.get("description", "")
            techs = project.get("technologies", [])
            story.append(Paragraph(f"<b>{name}</b>", styles["bold_body"]))
            if desc:
                story.append(Paragraph(desc, styles["body"]))
            if techs:
                story.append(Paragraph(
                    f"<i>Technologies: {', '.join(techs)}</i>",
                    styles["body"]
                ))
            story.append(Spacer(1, 4))

    # ── EDUCATION ───────────────────────────────────────────────
    education = resume_data.get("education", [])
    if education:
        add_section("EDUCATION")
        for edu in education:
            degree = edu.get("degree", "")
            institution = edu.get("institution", "")
            year = edu.get("year", "")
            text = f"<b>{degree}</b> — {institution}"
            if year:
                text += f"  <font color='#888888'>{year}</font>"
            story.append(Paragraph(text, styles["body"]))

    # ── CERTIFICATIONS ──────────────────────────────────────────
    certs = resume_data.get("certifications", [])
    if certs:
        add_section("CERTIFICATIONS")
        for cert in certs:
            if cert:
                story.append(Paragraph(f"• {cert}", styles["bullet"]))

    doc.build(story)
    logger.info(f"PDF saved to: {output_path}")
    return output_path
